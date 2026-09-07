"""
PPT Generation Orchestrator.

Generates PowerPoint presentations with:
- Background images (template mode or API mode)
- Editable text overlays
- Multiple layout styles
- Error handling and retry logic
"""

import json
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

# Import local modules
sys.path.insert(0, os.path.dirname(__file__))
from background import generate_background, get_available_styles
from layout import resolve_layout, auto_assign_layouts, get_recommended_layout
from text_render import render_slide_text, add_card_background


# ─── Constants ────────────────────────────────────────────────────────────────

SCRIPT_DIR = Path(__file__).parent
STYLES_DIR = SCRIPT_DIR / "styles"

MAX_RETRIES = 3
RETRY_DELAY = 2  # seconds

# Slide dimensions
SLIDE_WIDTH_16_9 = Inches(13.333)
SLIDE_HEIGHT_16_9 = Inches(7.5)
SLIDE_WIDTH_4_3 = Inches(10)
SLIDE_HEIGHT_4_3 = Inches(7.5)

# Background image dimensions (for template mode)
BG_WIDTH = 1920
BG_HEIGHT = 1080


# ─── Helper Functions ─────────────────────────────────────────────────────────

def load_plan(plan_file: str) -> dict:
    """Load and validate presentation plan."""
    with open(plan_file, "r", encoding="utf-8") as f:
        plan = json.load(f)

    # Validate required fields
    if "slides" not in plan:
        raise ValueError("Plan must contain 'slides' array")
    if not plan["slides"]:
        raise ValueError("Plan must contain at least one slide")

    return plan


def get_slide_dimensions(aspect_ratio: str) -> tuple:
    """Get slide dimensions in inches based on aspect ratio."""
    if aspect_ratio == "4:3":
        return SLIDE_WIDTH_4_3, SLIDE_HEIGHT_4_3
    return SLIDE_WIDTH_16_9, SLIDE_HEIGHT_16_9


def get_background_dimensions(aspect_ratio: str) -> tuple:
    """Get background image dimensions in pixels."""
    if aspect_ratio == "4:3":
        return 1600, 1200
    return BG_WIDTH, BG_HEIGHT


def generate_slide_background(
    slide_data: dict,
    style_config: dict,
    output_path: str,
    aspect_ratio: str = "16:9",
    use_api: bool = False,
) -> str:
    """
    Generate background for a single slide with retry logic.

    Returns:
        Path to generated background image
    """
    bg_width, bg_height = get_background_dimensions(aspect_ratio)

    for attempt in range(MAX_RETRIES):
        try:
            result = generate_background(
                style_config,
                output_path,
                width=bg_width,
                height=bg_height,
                use_api=use_api,
            )
            # Verify file was created
            if os.path.exists(result) and os.path.getsize(result) > 0:
                return result
        except Exception as e:
            if attempt < MAX_RETRIES - 1:
                time.sleep(RETRY_DELAY * (attempt + 1))
                continue
            raise Exception(f"Failed to generate background after {MAX_RETRIES} attempts: {e}")

    return None


def compose_slide(
    prs: Presentation,
    slide_data: dict,
    bg_path: str,
    style_config: dict,
    layout_name: str,
):
    """
    Compose a single slide with background image and text overlays.
    """
    # Get blank layout
    blank_layout = prs.slide_layouts[6]  # Blank layout

    # Add slide
    slide = prs.slides.add_slide(blank_layout)

    # Add background image
    if bg_path and os.path.exists(bg_path):
        with Image.open(bg_path) as img:
            # Convert to RGB if necessary
            if img.mode in ("RGBA", "P"):
                img = img.convert("RGB")

            # Save to bytes
            img_bytes = BytesIO()
            img.save(img_bytes, format="JPEG", quality=95)
            img_bytes.seek(0)

            # Add as background (full slide)
            slide.shapes.add_picture(
                img_bytes,
                Inches(0), Inches(0),
                prs.slide_width, prs.slide_height
            )

    # Add card background if enabled
    card_config = style_config.get("card", {})
    if card_config.get("enabled"):
        # Find the main content region for card placement
        regions = resolve_layout(layout_name, style_config,
                                 prs.slide_width / 914400,
                                 prs.slide_height / 914400)
        # Use first region as card reference
        if regions:
            first_region = list(regions.values())[0]
            add_card_background(
                slide,
                first_region.x - 0.02,
                first_region.y - 0.02,
                first_region.width + 0.04,
                first_region.height + 0.04,
                bg_color=card_config.get("bg"),
                border_color=card_config.get("border_color"),
                border_width=card_config.get("border_width", 1),
            )

    # Resolve layout and render text
    regions = resolve_layout(layout_name, style_config,
                             prs.slide_width / 914400,
                             prs.slide_height / 914400)

    render_slide_text(slide, slide_data, style_config, regions)

    # Add speaker notes
    notes = []
    if slide_data.get("title"):
        notes.append(f"Title: {slide_data['title']}")
    if slide_data.get("subtitle"):
        notes.append(f"Subtitle: {slide_data['subtitle']}")
    if slide_data.get("key_points"):
        notes.append("Key Points:")
        for point in slide_data["key_points"]:
            notes.append(f"  • {point}")

    if notes:
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        if text_frame is not None:
            text_frame.text = "\n".join(notes)


# ─── Main Generation Functions ────────────────────────────────────────────────

def generate_presentation(
    plan_file: str,
    output_file: str,
    style_name: str = None,
    use_api: bool = False,
    parallel: bool = True,
    preview_only: int = None,
) -> str:
    """
    Generate a complete PowerPoint presentation.

    Args:
        plan_file: Path to JSON presentation plan
        output_file: Path to output PPTX file
        style_name: Style name (overrides plan's style field)
        use_api: Whether to use Gemini API for backgrounds
        parallel: Whether to generate backgrounds in parallel
        preview_only: If set, only generate first N slides

    Returns:
        Status message
    """
    # Load plan
    plan = load_plan(plan_file)
    slides = plan.get("slides", [])

    # Determine style
    style_name = style_name or plan.get("style", "glassmorphism")
    styles = get_available_styles(str(STYLES_DIR))

    if style_name not in styles:
        raise ValueError(
            f"Style '{style_name}' not found. Available: {list(styles.keys())}"
        )

    style_config = styles[style_name]

    # Auto-assign layouts if not specified
    plan = auto_assign_layouts(plan)

    # Limit slides if preview mode
    if preview_only:
        slides = slides[:preview_only]

    # Get dimensions
    aspect_ratio = plan.get("aspect_ratio", "16:9")
    slide_width, slide_height = get_slide_dimensions(aspect_ratio)

    # Create presentation
    prs = Presentation()
    prs.slide_width = slide_width
    prs.slide_height = slide_height

    # Setup output directories
    output_dir = os.path.dirname(output_file) or "."
    bg_dir = os.path.join(output_dir, "_backgrounds")
    os.makedirs(bg_dir, exist_ok=True)

    # Generate backgrounds
    print(f"Generating {len(slides)} slides with style: {style_name}")

    def generate_slide_bg(slide_data):
        slide_num = slide_data.get("slide_number", 0)
        bg_path = os.path.join(bg_dir, f"slide-{slide_num:02d}.jpg")
        return slide_num, generate_slide_background(
            slide_data, style_config, bg_path, aspect_ratio, use_api
        )

    if parallel and len(slides) > 1:
        # Parallel generation
        bg_results = {}
        with ThreadPoolExecutor(max_workers=min(4, len(slides))) as executor:
            futures = {executor.submit(generate_slide_bg, s): s for s in slides}
            for future in as_completed(futures):
                slide_num, bg_path = future.result()
                bg_results[slide_num] = bg_path
                print(f"  ✓ Slide {slide_num} background ready")
    else:
        # Sequential generation
        bg_results = {}
        for slide_data in slides:
            slide_num = slide_data.get("slide_number", 0)
            bg_path = os.path.join(bg_dir, f"slide-{slide_num:02d}.jpg")
            result = generate_slide_background(
                slide_data, style_config, bg_path, aspect_ratio, use_api
            )
            bg_results[slide_num] = result
            print(f"  ✓ Slide {slide_num} background ready")

    # Compose slides
    print("Composing slides...")
    for slide_data in slides:
        slide_num = slide_data.get("slide_number", 0)
        bg_path = bg_results.get(slide_num)
        layout_name = slide_data.get("layout", "content")

        compose_slide(prs, slide_data, bg_path, style_config, layout_name)
        print(f"  ✓ Slide {slide_num} composed")

    # Save presentation
    prs.save(output_file)
    print(f"\n✓ Presentation saved to: {output_file}")

    # Cleanup backgrounds (keep them for debugging)
    # import shutil
    # shutil.rmtree(bg_dir, ignore_errors=True)

    return f"Successfully generated {len(slides)}-slide presentation to {output_file}"


# ─── CLI Interface ────────────────────────────────────────────────────────────

def main():
    import argparse

    parser = argparse.ArgumentParser(
        description="Generate PowerPoint presentation from plan"
    )
    parser.add_argument(
        "--plan-file",
        required=True,
        help="Path to JSON presentation plan",
    )
    parser.add_argument(
        "--output-file",
        required=True,
        help="Output path for PPTX file",
    )
    parser.add_argument(
        "--style",
        help="Style name (overrides plan's style)",
        choices=["glassmorphism", "dark-premium", "gradient-modern", "minimal-light", "vibrant"],
    )
    parser.add_argument(
        "--use-api",
        action="store_true",
        help="Use Gemini API for background generation",
    )
    parser.add_argument(
        "--no-parallel",
        action="store_true",
        help="Disable parallel background generation",
    )
    parser.add_argument(
        "--preview",
        type=int,
        help="Only generate first N slides (for preview)",
    )
    parser.add_argument(
        "--list-styles",
        action="store_true",
        help="List available styles and exit",
    )

    args = parser.parse_args()

    # List styles mode
    if args.list_styles:
        styles = get_available_styles(str(STYLES_DIR))
        print("Available styles:")
        for name, style in styles.items():
            print(f"  {name}: {style['description']}")
            print(f"    Best for: {style['best_for']}")
        return

    try:
        result = generate_presentation(
            args.plan_file,
            args.output_file,
            style_name=args.style,
            use_api=args.use_api,
            parallel=not args.no_parallel,
            preview_only=args.preview,
        )
        print(result)
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
